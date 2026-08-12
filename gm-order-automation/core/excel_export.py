from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.settings import EXPORTS_DIR
from core.models import Document

HEADERS = [
    ("doc_no", "문서번호"),
    ("partner", "거래처"),
    ("item", "품번"),
    ("item_name", "품명"),
    ("qty", "수량"),
    ("order_date", "발주일"),
    ("due_date", "납기일"),
    ("status", "상태"),
    ("source_file", "원본파일"),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
FAIL_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")


def export_documents(documents: list[Document], output_path: str | Path | None = None) -> Path:
    if output_path is None:
        output_path = EXPORTS_DIR / f"발주현황_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "문서별 처리현황"

    for col_idx, (_, label) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_idx, doc in enumerate(documents, start=2):
        for col_idx, (field_name, _) in enumerate(HEADERS, start=1):
            value = getattr(doc, field_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if field_name == "status" and value not in ("성공", "완료", "정상", ""):
                cell.fill = FAIL_FILL

    for col_idx, (field_name, label) in enumerate(HEADERS, start=1):
        width = max(len(label) + 4, 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    summary = wb.create_sheet("요약")
    summary["A1"] = "총 문서 수"
    summary["B1"] = len(documents)
    summary["A2"] = "생성 시각"
    summary["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    summary.column_dimensions["A"].width = 16

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
